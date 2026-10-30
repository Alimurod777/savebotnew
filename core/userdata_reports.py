"""Offline metadata reports. No Telegram or database access."""
from __future__ import annotations

import html
import base64
import hashlib
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import Enum
from html.parser import HTMLParser
from itertools import zip_longest
from pathlib import Path
from types import SimpleNamespace
from typing import Dict, Optional, Tuple, Union

CATEGORY_LABELS = {
    'private_channels': 'Private Channels', 'private_groups': 'Private Groups',
    'bots': 'Bots', 'private_chats': 'Private Chats', 'public_channels': 'Public Channels',
    'public_groups': 'Public Groups', 'unclassified': 'Unclassified',
}
CATEGORY_IDS = tuple(CATEGORY_LABELS)
CATEGORIES = tuple(CATEGORY_LABELS.values())  # Workbook names; never filter identifiers.
PRIVATE_CATEGORIES = ('private_channels', 'private_groups')
COMPACT_CATEGORIES = ('bots', 'private_chats', 'public_channels', 'public_groups')
NOTE = ('Metadata only; no message contents. Username confirms public discoverability; '
        'missing usernames do not prove privacy. Unknown fields may be unavailable or unsupported. '
        'Private user chats describe the dialog type, not profile visibility. '
        'Detailed inspection is limited to non-public channels/groups; public-group topics are not requested. '
        'Topic lists are accessible metadata and may be incomplete; an empty list is not proof of no topics. '
        'Compact XLSX sheets contain name, ID, username and type. Detailed sheets preserve long values in continuation columns. '
        'XLSX omits unsupported control characters; values are never silently truncated.')


def clean(value):
    if isinstance(value, Enum):
        value = value.value
    return str(value).encode('utf-8', 'replace').decode('utf-8') if value is not None else ''


def type_name(value):
    return clean(getattr(value, 'value', value)).split('.')[-1].lower()


def chat_kind(chat_type, is_bot=False, fallback='unknown'):
    kind = type_name(chat_type)
    if kind == 'bot' or (kind == 'private' and is_bot):
        return 'bot'
    if kind in ('group', 'supergroup', 'forum', 'monoforum', 'gigagroup'):
        return 'group'
    if kind in ('channel', 'private'):
        return kind
    return type_name(fallback) if not kind or kind == 'unknown' else 'unknown'


def active_username(chat):
    username = getattr(chat, 'username', None)
    if username:
        return clean(username)
    for alias in getattr(chat, 'usernames', None) or ():
        if getattr(alias, 'active', False) and getattr(alias, 'username', None):
            return clean(alias.username)
    return ''


def privacy_evidence(chat, kind=None):
    """Use positive dialog/type evidence or provenance explicitly set by collection."""
    kind = kind or chat_kind(getattr(chat, 'type', getattr(chat, 'chat_type', '')),
                             getattr(chat, 'is_bot', False), getattr(chat, 'category', 'unknown'))
    if kind == 'private':
        return 'PRIVATE', 'Telegram private dialog type'
    if kind == 'group' and type_name(getattr(chat, 'type', getattr(chat, 'chat_type', ''))) == 'group':
        return 'PRIVATE', 'Telegram basic group type'
    if kind in ('channel', 'group'):
        visibility = getattr(chat, 'visibility', 'UNKNOWN')
        verified = getattr(chat, 'visibility_verified', False) is True
        source = getattr(chat, 'visibility_source', '')
        if source == 'Conflicting privacy evidence':
            return 'UNKNOWN', source
        if verified and source and visibility in ('PUBLIC', 'PRIVATE'):
            return visibility, clean(source)
        if active_username(chat):
            return 'PUBLIC', 'Telegram active username'
    return 'UNKNOWN', 'Unverified / unavailable'


def category_id(kind, privacy):
    if kind == 'bot':
        return 'bots'
    if kind == 'private':
        return 'private_chats'
    if kind in ('channel', 'group') and privacy in ('PUBLIC', 'PRIVATE'):
        return privacy.lower() + ('_channels' if kind == 'channel' else '_groups')
    return 'unclassified'


DETAIL_FIELDS = (
    ('description', 'Description'), ('member_count', 'Member count'),
    ('visibility', 'Collected privacy status'), ('visibility_source', 'Collected privacy source'),
    ('visibility_verified', 'Privacy verified'),
    ('latest_message_id', 'Latest message ID'), ('is_bot', 'Is bot'),
    ('has_discussion', 'Has discussion'), ('discussion_group_id', 'Discussion group ID'),
    ('discussion_group_title', 'Discussion group title'), ('latest_public_link', 'Latest public link'),
    ('comment_availability', 'Comment availability'), ('inspection_status', 'Inspection'),
    ('error_reason', 'Inspection error'), ('forum_status', 'Forum metadata status'),
)
TOPIC_FIELDS = (
    ('title', 'Title'), ('top_message', 'Top message ID'), ('closed', 'Closed'),
    ('hidden', 'Hidden'), ('first_message_id', 'First message ID'),
    ('latest_message_id', 'Latest message ID'), ('inspection_status', 'Inspection'),
)


def _topic_text(topics):
    grouped = {}
    for topic in topics:
        values = grouped.setdefault(str(topic.topic_id), {})
        for attr, label in TOPIC_FIELDS:
            value = getattr(topic, attr, None)
            if value is not None:
                values.setdefault(label, {})
                text = clean(value)
                values[label][text] = None
    return '\n'.join(f"Topic {topic_id}\n" + '\n'.join(f"  {label}: {' | '.join(values)}" for label, values in fields.items())
                     for topic_id, fields in grouped.items())


@dataclass(frozen=True)
class Record:
    name: str
    chat_id: str
    username: str
    chat_type: str
    privacy: str
    source: str
    forum: str
    topics: str
    metadata: str
    category: str


@dataclass(frozen=True)
class Snapshot:
    records: Tuple[Record, ...]
    timestamp: str
    discovered: int
    note: str
    target_user_id: Optional[int] = None

    @property
    def counts(self) -> Dict[str, int]:
        counts = dict.fromkeys(CATEGORY_IDS, 0)
        for record in self.records:
            counts[record.category] += 1
        return counts


def normalize(job) -> Snapshot:
    # Merge complementary duplicate observations, retaining conflicting values as
    # metadata rather than dropping the later record wholesale.
    merged, observations = {}, defaultdict(dict)
    for rec in job.records:
        canonical_type = type_name(rec.chat_type)
        key = (str(rec.chat_id), canonical_type or type_name(rec.category))
        if key not in merged:
            merged[key] = SimpleNamespace(**vars(rec))
            merged[key].topics = list(rec.topics)
            continue
        previous = merged[key]
        previous.topics.extend(rec.topics)
        if getattr(rec, 'visibility_verified', False):
            if (getattr(previous, 'visibility_verified', False)
                    and previous.visibility != rec.visibility):
                previous.visibility = 'UNKNOWN'
                previous.visibility_source = 'Conflicting privacy evidence'
                previous.visibility_verified = False
            elif getattr(previous, 'visibility_source', '') != 'Conflicting privacy evidence':
                previous.visibility, previous.visibility_source = rec.visibility, rec.visibility_source
                previous.visibility_verified = True
        for attr in ('title', 'username',
                     'is_forum', 'topic_count', 'topics_status', *(f[0] for f in DETAIL_FIELDS)):
            value, old = getattr(rec, attr, None), getattr(previous, attr, None)
            missing = old is None or old == '' or (attr in ('topics_status', 'forum_status', 'comment_availability')
                and old in ('UNKNOWN', 'Unknown / unverified', 'Unavailable / not enumerated'))
            if missing:
                setattr(previous, attr, value)
            elif value is not None and value != '' and value != old:
                note = f'Other observed {attr}: {clean(value)}'
                observations[key][note] = None
    records = []
    for key, rec in merged.items():
        kind = chat_kind(rec.chat_type, getattr(rec, 'is_bot', False), rec.category)
        privacy, source = privacy_evidence(rec, kind)
        topic_status = getattr(rec, 'topics_status', 'Unavailable / not enumerated')
        topic_count = getattr(rec, 'topic_count', None)
        topics = clean(topic_status) + ('; total unknown' if topic_count is None else f'; API reported count: {topic_count}')
        topic_text = _topic_text(rec.topics)
        if topic_text:
            topics += '\n' + topic_text
        metadata = '\n'.join(f'{label}: {clean(value)}' for label, value in (
            (label, getattr(rec, attr, None)) for attr, label in DETAIL_FIELDS) if value is not None)
        if observations[key]:
            metadata += '\n' + '\n'.join(observations[key])
        records.append(Record(clean(rec.title) or '(name unavailable)', str(rec.chat_id), clean(rec.username),
                              type_name(rec.chat_type), privacy, source,
                              getattr(rec, 'forum_status', 'Unknown / unverified') if rec.is_forum is None else ('Yes' if rec.is_forum else 'No'),
                              topics, metadata or 'Unavailable', category_id(kind, privacy)))
    records.sort(key=lambda r: CATEGORY_IDS.index(r.category))
    partial = sum(r.inspection_status != 'COMPLETE' for r in job.records)
    limitations = f' Collected {len(records)} of {job.dialogs_discovered} discovered chats. Operation failures: {len(job.failures)}. Partial inspections: {partial}.'
    if any(f.get('stage') == 'scan' for f in job.failures):
        limitations += ' Dialog scan interrupted or configured chat limit reached; discovered count is a lower bound.'
    return Snapshot(tuple(records), datetime.now(timezone.utc).isoformat(), job.dialogs_discovered,
                    NOTE + limitations, job.target_user_id)


def render_html(snapshot: Snapshot, path: Union[str, Path]) -> None:
    esc = lambda v: html.escape(clean(v), quote=True)
    cards = []
    for n, r in enumerate(snapshot.records, 1):
        link = ''
        if re.fullmatch(r'[A-Za-z0-9_]{1,64}', r.username):
            link = f'<a href="https://t.me/{esc(r.username)}" rel="noreferrer noopener" target="_blank">@{esc(r.username)}</a>'
        layout = 'detailed' if r.category in PRIVATE_CATEGORIES else 'compact'
        cards.append(f'''<article class="{layout}" data-chat-id="{esc(r.chat_id)}" data-chat-type="{esc(r.chat_type)}" data-category="{esc(r.category)}" data-privacy="{esc(r.privacy)}" data-search="{esc(r.name+' '+r.chat_id+' '+r.username)}">
<details><summary><small>#{n} · {esc(CATEGORY_LABELS[r.category])} · {esc(r.privacy)}</small><h2>{esc(r.name)}</h2><code>{esc(r.chat_id)}</code></summary>
<p>{link}</p><dl><dt>Full name</dt><dd>{esc(r.name)}</dd><dt>Type</dt><dd>{esc(r.chat_type)}</dd><dt>Privacy evidence</dt><dd>{esc(r.source)}</dd><dt>Forum enabled</dt><dd>{esc(r.forum)}</dd></dl>
<details><summary>Topic metadata</summary><pre>{esc(r.topics)}</pre></details><pre>{esc(r.metadata)}</pre></details>
<button class="copy" data-id="{esc(r.chat_id)}">Copy ID</button></article>''')
    options = ''.join(f'<option value="{cat}">{esc(label)}</option>' for cat, label in CATEGORY_LABELS.items())
    stats = ''.join(f'<button class="stat" data-filter="{cat}"><strong>{count}</strong>{esc(CATEGORY_LABELS[cat])}</button>' for cat, count in snapshot.counts.items())
    content = '''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; script-src SCRIPT_POLICY; style-src 'unsafe-inline'; connect-src 'none'; img-src 'none'; base-uri 'none'; form-action 'none'; object-src 'none'">
<title>Userdata inventory</title><style>
*{box-sizing:border-box}body{margin:0;background:#0c1220;color:#edf3ff;font:16px system-ui,sans-serif}main{max-width:1200px;margin:auto;padding:16px}h1{font-size:clamp(24px,5vw,38px)}h2{font-size:18px;margin:10px 0;overflow-wrap:anywhere}small,dt{color:#a8bad4}header,p{line-height:1.6}.stats,.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,270px),1fr));gap:12px}.stats div,article{background:#172238;padding:16px;border:1px solid #304362;border-radius:14px;min-width:0}.stats strong{display:block;font-size:26px;color:#69e0cc}nav{position:sticky;top:0;background:#0c1220;padding:12px 0;display:flex;gap:8px;flex-wrap:wrap}input,select,button{font:inherit;background:#23324d;color:#fff;border:1px solid #526584;border-radius:8px;padding:10px;max-width:100%}input{flex:1;min-width:160px}button,summary{cursor:pointer}button{margin-top:8px}summary{overflow-wrap:anywhere}code,pre,dd{overflow-wrap:anywhere;white-space:pre-wrap;margin-left:0}a{color:#79ceff}[hidden]{display:none!important}input:focus,button:focus,select:focus,summary:focus{outline:2px solid #69e0cc}
</style><main><header><small>TELEGRAM INVENTORY · OFFLINE REPORT</small><h1>Telegram inventory</h1>'''
    content += '''<style>h2{display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden}article>details[open] h2{display:block}.stat{text-align:left;margin:0;min-height:100px}.stat strong{display:block;font-size:28px;color:#69e0cc}.stat:nth-child(-n+2){border-color:#69e0cc;background:#163437}.detailed{border-left:3px solid #69e0cc}.grid{align-items:start}.grid>article{content-visibility:auto;contain-intrinsic-size:auto 200px}.compact{grid-column:1/-1;display:grid;grid-template-columns:minmax(0,1fr) auto;gap:12px;padding:12px 16px}.compact summary{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,220px);gap:4px 12px}.compact small{grid-column:1/-1}.compact h2{font-size:16px;margin:0;-webkit-line-clamp:1}.compact .copy{align-self:start}.compact details{min-width:0}nav{z-index:1}nav input{min-width:min(100%,160px)}@media(max-width:480px){main{padding:12px}.stats{grid-template-columns:1fr 1fr}.stat{font-size:13px;padding:12px}nav select{flex:1;min-width:0}nav input{flex-basis:100%}.compact summary{grid-template-columns:minmax(0,1fr)}.compact{gap:8px;padding:12px}.compact .copy{font-size:13px}}@media(prefers-reduced-motion:reduce){*{scroll-behavior:auto}}</style>'''
    account = f'<p>Account: {esc(snapshot.target_user_id)}</p>' if snapshot.target_user_id is not None else ''
    content += f'{account}<p>{esc(snapshot.timestamp)} · UTC<br>{len(snapshot.records)} chats · {snapshot.discovered} discovered</p><p>{esc(snapshot.note)}</p></header><section class="stats" aria-label="Category totals">{stats}</section>'
    content += f'''<nav aria-label="Report filters"><input id="search" aria-label="Search name, ID or username" placeholder="Search name, ID, username"><select id="category" aria-label="Category"><option value="">All categories</option>{options}</select><select id="privacy" aria-label="Privacy"><option value="">All privacy states</option><option>PRIVATE</option><option>PUBLIC</option><option>UNKNOWN</option></select><button id="clear">Clear search</button></nav><p id="count" aria-live="polite"></p><p id="empty" hidden>No chats match the current category, privacy and search filters.</p><section class="grid">{''.join(cards)}</section></main>'''
    content += '''<script>
const cards=Array.from(document.querySelectorAll('article')),search=document.getElementById('search'),category=document.getElementById('category'),privacy=document.getElementById('privacy');
const fold=s=>s.normalize('NFKC').toLowerCase().replace(/ß/g,'ss').replace(/ς/g,'σ');
const index=cards.map(c=>fold(c.dataset.search));
function filter(){const q=fold(search.value.trim());let n=0;cards.forEach((c,i)=>{const show=index[i].includes(q)&&(!category.value||c.dataset.category===category.value)&&(!privacy.value||c.dataset.privacy===privacy.value);if(c.hidden!==!show)c.hidden=!show;if(show)n++;});document.getElementById('count').textContent=n+' of '+cards.length+' chats';document.getElementById('empty').hidden=n!==0;}
search.addEventListener('input',filter);category.addEventListener('change',filter);privacy.addEventListener('change',filter);document.getElementById('clear').onclick=()=>{search.value='';filter();};
document.querySelectorAll('.stat').forEach(b=>b.onclick=()=>{category.value=b.dataset.filter;filter();});
document.querySelectorAll('.copy').forEach(b=>b.onclick=async()=>{try{await navigator.clipboard.writeText(b.dataset.id);b.textContent='Copied';}catch(e){const t=document.createElement('textarea');t.value=b.dataset.id;document.body.appendChild(t);t.select();b.textContent=document.execCommand('copy')?'Copied':'Select ID above to copy';t.remove();}});filter();</script></html>'''
    script = content.rsplit('<script>', 1)[1].split('</script>', 1)[0]
    digest = base64.b64encode(hashlib.sha256(script.encode('utf-8')).digest()).decode('ascii')
    content = content.replace('script-src SCRIPT_POLICY', f"script-src 'sha256-{digest}'", 1)
    Path(path).write_text(content, encoding='utf-8')
    validate_html(snapshot, path)


class ReportValidationError(ValueError):
    """Safe operational error: never include inventory values in its message."""


class ReportLimitError(ValueError):
    """A report cannot represent a field within its format's limits."""


def validate_html(snapshot: Snapshot, path: Union[str, Path]) -> None:
    class InventoryParser(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.records = Counter()

        def handle_starttag(self, tag, attrs):
            if tag == 'article':
                attrs = dict(attrs)
                self.records[(attrs.get('data-chat-id'), attrs.get('data-chat-type'), attrs.get('data-category'))] += 1

    parser = InventoryParser()
    parser.feed(Path(path).read_text(encoding='utf-8'))
    expected = Counter((r.chat_id, r.chat_type, r.category) for r in snapshot.records)
    if parser.records != expected:
        raise ReportValidationError('HTML inventory identities or category counts do not match the shared snapshot')


DETAILED_HEADERS = ('Number', 'Chat Name', 'Chat ID', 'Username', 'Chat Type', 'Privacy Status',
                    'Forum Enabled', 'Topic Information', 'Additional Available Metadata')
COMPACT_HEADERS = ('Name', 'Chat ID', 'Username', 'Chat Type')


def _xlsx_values(record, number):
    if record.category in COMPACT_CATEGORIES:
        return [record.name, record.chat_id, record.username, record.chat_type]
    return [number, record.name, record.chat_id, record.username, record.chat_type,
            record.privacy + ' — ' + record.source, record.forum, record.topics, record.metadata]


def _excel_value(value):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub('', clean(value)) if isinstance(value, str) else value


def _excel_parts(value):
    """Lossless UTF-16 cell chunks, including whitespace and surrogate pairs."""
    if not isinstance(value, str):
        return [value]
    data, chunks, start = value.encode('utf-16-le'), [], 0
    while start < len(data):
        end = min(start + 32767 * 2, len(data))
        if end < len(data) and 0xD800 <= int.from_bytes(data[end-2:end], 'little') <= 0xDBFF:
            end -= 2
        chunks.append(data[start:end].decode('utf-16-le'))
        start = end
    return chunks or ['']


def _rows_by_category(snapshot):
    grouped = {cat: [] for cat in CATEGORY_IDS}
    for record in snapshot.records:
        grouped[record.category].append(record)
    return grouped


def validate_xlsx(snapshot: Snapshot, path: Union[str, Path]) -> None:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ['Summary', *CATEGORIES]:
            raise ReportValidationError('XLSX worksheet structure does not match the shared snapshot')
        summary = {row[0]: row[1] for row in book['Summary'].values}
        if summary.get('Total chats') != len(snapshot.records):
            raise ReportValidationError('XLSX total does not match the shared snapshot')
        grouped = _rows_by_category(snapshot)
        for cat, label in CATEGORY_LABELS.items():
            sheet = book[label]
            headers = COMPACT_HEADERS if cat in COMPACT_CATEGORIES else DETAILED_HEADERS
            actual_headers = next(sheet.values)
            if tuple(actual_headers[:len(headers)]) != headers or summary.get(label) != len(grouped[cat]):
                raise ReportValidationError('XLSX category headers or counts do not match the shared snapshot')
            for number, (row, record) in enumerate(zip_longest(sheet.iter_rows(min_row=2), grouped[cat]), 1):
                if row is None or record is None:
                    raise ReportValidationError('XLSX category rows do not match the shared snapshot')
                if any(cell.data_type == 'f' for cell in row):
                    raise ReportValidationError('XLSX inventory must not contain formulas')
                expected = _xlsx_values(record, number)
                for index, value in enumerate(expected):
                    actual = row[index].value
                    if isinstance(value, str):
                        actual = '' if actual is None else actual
                        part = 2
                        while f'{headers[index]} (continued {part})' in actual_headers:
                            extra_index = actual_headers.index(f'{headers[index]} (continued {part})')
                            actual += row[extra_index].value or ''
                            part += 1
                    if actual != _excel_value(value):
                        raise ReportValidationError('XLSX inventory values do not match the shared snapshot')
    finally:
        book.close()


def render_xlsx(snapshot: Snapshot, path: Union[str, Path]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    def append(sheet, values):
        # Validate before openpyxl's cell binder silently truncates long strings.
        safe = []
        for value in values:
            if isinstance(value, str):
                value = _excel_value(value)
                if len(_excel_parts(value)) > 1:
                    raise ReportLimitError('Excel cell limit exceeded; report not truncated')
            safe.append(value)
        sheet.append(safe)
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    append(summary, ['Telegram metadata inventory', snapshot.timestamp])
    summary.append(['Total chats', len(snapshot.records)])
    summary.append(['Discovered chats', snapshot.discovered])
    append(summary, ['Limitations', snapshot.note])
    for cat, count in snapshot.counts.items():
        label = CATEGORY_LABELS[cat]
        summary.append([label, count])
        summary.cell(summary.max_row, 1).hyperlink = f"#'{label}'!A1"
    if snapshot.target_user_id is not None:
        append(summary, ['Account ID', str(snapshot.target_user_id)])
    grouped = _rows_by_category(snapshot)
    for cat, label in CATEGORY_LABELS.items():
        sheet = book.create_sheet(label)
        compact = cat in COMPACT_CATEGORIES
        headers = COMPACT_HEADERS if compact else DETAILED_HEADERS
        sheet.append(headers)
        continuations = {}
        for n, r in enumerate(grouped[cat], 1):
            values = _xlsx_values(r, n)
            row_number = n + 1
            for index, value in enumerate(values, 1):
                parts = _excel_parts(_excel_value(value))
                if compact and len(parts) > 1:
                    raise ReportLimitError('Excel compact field exceeds cell limit; report not truncated')
                sheet.cell(row_number, index, parts[0])
                for part_number, part in enumerate(parts[1:], 2):
                    key = (index, part_number)
                    if key not in continuations:
                        continuations[key] = len(headers) + len(continuations) + 1
                        if continuations[key] > 16384:
                            raise ReportLimitError('Excel column limit exceeded; report not truncated')
                        sheet.cell(1, continuations[key], f'{headers[index-1]} (continued {part_number})')
                    sheet.cell(row_number, continuations[key], part)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = 'A2'
        widths = (40, 25, 25, 24) if compact else (10, 32, 25, 24, 24, 40, 26, 50, 50)
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths[index-1] if index <= len(widths) else 50
        for cell in sheet['B' if compact else 'C'][1:]:
            cell.number_format = '@'
    summary.column_dimensions['A'].width = 36
    summary.column_dimensions['B'].width = 85
    for sheet in book:
        sheet.freeze_panes = 'A2'
        for row in sheet:
            sheet.row_dimensions[row[0].row].height = 48 if row[0].row > 1 else 30
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _excel_value(cell.value)
                    cell.data_type = 's'  # Explicit strings prevent formulas, including = + - @.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11, color='FFFFFF' if cell.row == 1 else '182C45', bold=cell.row == 1)
                cell.fill = PatternFill('solid', fgColor='203A5C' if cell.row == 1 else ('EDF3FA' if cell.row % 2 == 0 else 'FFFFFF'))
    summary.row_dimensions[4].height = 120
    book.save(path)
    validate_xlsx(snapshot, path)
