"""Inventory transformation regressions: synthetic data, no account credentials."""
from __future__ import annotations

from dataclasses import replace
from html.parser import HTMLParser
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace as NS
import json
import re
import subprocess
import unittest
from unittest.mock import AsyncMock, patch

from core.userdata_reports import (
    CATEGORY_IDS, CATEGORY_LABELS, CATEGORIES, COMPACT_CATEGORIES, COMPACT_HEADERS,
    DETAILED_HEADERS, ReportValidationError, normalize, render_html, render_xlsx,
    validate_html, validate_xlsx,
)
from TechVJ import userdata_command as command


def fixture_records():
    D, T = command.DialogRecord, command.TopicRecord
    verified = dict(visibility='PRIVATE', visibility_verified=True,
                    visibility_source='Verified full Channel metadata')
    return [
        D(-100000000000001, 'channel', title='Ўзбек Straße 😀 <script>bad()</script> '+ 'x'*700,
          chat_type='ChatType.CHANNEL', description='channel description', member_count=0,
          latest_message_id=123, has_discussion=True, discussion_group_id=-100000000000002,
          discussion_group_title='Discussion <&>', comment_availability='UNKNOWN', **verified),
        D(-100000000000002, 'unknown', title='Private forum', chat_type='ChatType.FORUM',
          description='group description', member_count=0, is_forum=True, forum_status='Known',
          topic_count=1, topics_status='Complete accessible topic list', latest_message_id=999,
          has_discussion=False, inspection_status='PARTIAL', error_reason='Metadata unavailable',
          topics=[T(5, 'Topic <&> 😀', top_message=55, first_message_id=5,
                    latest_message_id=55, closed=False, hidden=True)], **verified),
        D(3, 'bot', title='=HYPERLINK("evil")', chat_type='ChatType.BOT', username='helper_bot', is_bot=True),
        D(4, 'private', title='Person', chat_type='private'),
        D(-100000000000005, 'channel', title='Public channel', chat_type='channel', username='public_channel'),
        D(-100000000000006, 'group', title='Public group', chat_type='supergroup', username='public_group'),
        D(-100000000000007, 'channel', title='Unknown', chat_type='channel', visibility='PRIVATE'),
    ]


def snapshot(records=None):
    rows = fixture_records() if records is None else records
    return normalize(command.InventoryJob('test', 42, records=rows, dialogs_discovered=len(rows)))


def raw_channel(**values):
    fields = dict(id=7, title='Synthetic', photo=command.types.ChatPhotoEmpty(), date=1,
                  broadcast=True, access_hash=9)
    fields.update(values)
    return command.types.Channel(**fields)


class Markup(HTMLParser):
    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.articles, self.stats, self.options = [], [], []
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'article': self.articles.append(attrs)
        if tag == 'button' and attrs.get('class') == 'stat': self.stats.append(attrs)
        if tag == 'option': self.options.append(attrs)


class Classification(unittest.TestCase):
    def test_all_categories_and_enum_variants(self):
        s = snapshot()
        self.assertEqual(tuple(r.category for r in s.records), CATEGORY_IDS)
        self.assertEqual(list(s.counts.values()), [1]*7)
        for value in (command.ChatType.FORUM, 'ChatType.FORUM', 'forum', 'ChatType.MONOFORUM',
                      command.ChatType.GROUP, 'supergroup'):
            self.assertEqual(command._classify_dialog(NS(chat=NS(type=value))), 'group')

    def test_verified_private_channels_and_supergroups(self):
        for options in ({'broadcast':True}, {'broadcast':False,'megagroup':True,'forum':True}):
            with self.subTest(options=options):
                privacy, source = command._raw_channel_privacy(raw_channel(**options))
                self.assertEqual(privacy,'PRIVATE')
                self.assertIn('non-min',source)

    def test_missing_username_min_geo_left_and_unsupported_stay_unknown(self):
        for raw in (NS(username=None), raw_channel(min=True), raw_channel(has_geo=True),
                    raw_channel(left=True), raw_channel(monoforum=True)):
            self.assertEqual(command._raw_channel_privacy(raw)[0],'UNKNOWN')
        self.assertEqual(snapshot().records[-1].category,'unclassified')

    def test_active_and_inactive_aliases(self):
        alias=command.types.Username(username='alias',active=True)
        self.assertEqual(command._raw_channel_privacy(raw_channel(usernames=[alias]))[0],'PUBLIC')
        self.assertEqual(command._visibility(NS(type=command.ChatType.FORUM,username=None,usernames=[alias])),'PUBLIC')
        alias=command.types.Username(username='inactive',active=False)
        self.assertEqual(command._raw_channel_privacy(raw_channel(usernames=[alias]))[0],'PRIVATE')

    def test_basic_group_has_positive_type_evidence(self):
        s=snapshot([command.DialogRecord(-7,'group',title='Basic',chat_type=command.ChatType.GROUP)])
        self.assertEqual(s.records[0].category,'private_groups')
        self.assertIn('basic group',s.records[0].source)

    def test_duplicates_merge_fields_and_preserve_distinct_types(self):
        first=fixture_records()[1]
        first.member_count=None
        second=replace(first,chat_type='forum',member_count=44,description='Other observed description',
                       topics=[command.TopicRecord(5,'Updated topic',latest_message_id=88)])
        other_type=replace(first,chat_type='channel',category='channel')
        s=snapshot([first,second,first,other_type])
        self.assertEqual(len(s.records),2)
        forum=next(r for r in s.records if r.chat_type=='forum')
        self.assertIn('Member count: 44',forum.metadata)
        self.assertIn('group description',forum.metadata)
        self.assertIn('Other observed description',forum.metadata)
        self.assertIn('Updated topic',forum.topics)
        self.assertIn('Latest message ID: 55 | 88',forum.topics)
        self.assertEqual(forum.topics.count('Topic 5\n'),1)

    def test_conflicting_verified_privacy_is_not_silently_overwritten(self):
        first=fixture_records()[0]
        second=replace(first,visibility='PUBLIC',visibility_source='Active alias observed')
        s=snapshot([first,second])
        self.assertEqual(s.records[0].category,'unclassified')
        self.assertEqual(s.records[0].privacy,'UNKNOWN')


class InventoryCollection(unittest.IsolatedAsyncioTestCase):
    async def test_basic_group_member_count_uses_existing_participant_metadata(self):
        client=AsyncMock();client.resolve_peer.return_value=NS(chat_id=7)
        members=command.types.ChatParticipants(chat_id=7,participants=[NS(),NS()],version=1)
        client.invoke.return_value=NS(full_chat=NS(id=7,participants=members),chats=[NS(id=7)])
        meta=await command._get_chat_metadata(client,NS(id=-7,type=command.ChatType.GROUP))
        self.assertEqual(meta.members_count,2)
        self.assertEqual(meta.visibility,'PRIVATE')
        self.assertTrue(meta.visibility_verified)
        client.get_chat_members.assert_not_called()

    async def test_unset_topic_protocol_flags_are_known_false(self):
        topic=NS(id=5,title='Topic',top_message=55,closed=None,hidden=None)
        client=AsyncMock();client.invoke.return_value=NS(count=1,topics=[topic])
        with patch.object(command.types,'ForumTopic',NS):
            result=await command._fetch_forum_topics(client,42)
        self.assertIs(result[0].closed,False)
        self.assertIs(result[0].hidden,False)

    async def test_missing_supplementary_fields_do_not_erase_known_dialog_values(self):
        full=NS(type=command.ChatType.FORUM,description=None,members_count=None,is_forum=None,
                username=None,visibility='UNKNOWN',visibility_verified=False,visibility_source='Unavailable')
        chat=NS(id=-1007,type=command.ChatType.FORUM,title='Forum',username=None,members_count=0)
        with patch.object(command,'_get_chat_metadata',AsyncMock(return_value=full)),patch.object(command,'_fetch_forum_topics',AsyncMock()) as topics:
            rec=await command._inspect_group(AsyncMock(),chat,NS(top_message=None),command.InventoryJob('test',42))
        self.assertEqual(rec.member_count,0)
        self.assertIs(rec.is_forum,True)
        topics.assert_not_awaited()

    async def test_public_groups_and_channels_do_not_request_full_metadata_or_topics(self):
        for method, kind in ((command._inspect_group, command.ChatType.FORUM),
                             (command._inspect_group, command.ChatType.SUPERGROUP),
                             (command._inspect_channel,command.ChatType.CHANNEL)):
            with self.subTest(kind=kind), patch.object(command,'_get_chat_metadata',AsyncMock()) as metadata, patch.object(command,'_fetch_forum_topics',AsyncMock()) as topics:
                chat=NS(id=-1007,type=kind,title='Public',username='public',members_count=42)
                result=await method(AsyncMock(),chat,NS(top_message=NS(id=11)),command.InventoryJob('test',42))
                metadata.assert_not_awaited();topics.assert_not_awaited()
                self.assertEqual(result.visibility,'PUBLIC')
                self.assertEqual(result.member_count,42)

    async def test_private_forum_metadata_and_topics_survive_collection_and_normalization(self):
        client=AsyncMock();client.resolve_peer.return_value=NS(channel_id=7,access_hash=9)
        client.invoke.return_value=NS(full_chat=NS(id=7,about='Private forum',participants_count=42),
                                      chats=[raw_channel(broadcast=False,megagroup=True,forum=True)])
        chat=NS(id=-1000000000007,type=command.ChatType.FORUM,title='Forum',username=None)
        job=command.InventoryJob('test',42)
        async def topics(user_client,chat_id,record):
            record.topics_status='Complete accessible topic list';record.topic_count=1
            return [command.TopicRecord(5,'Private topic',latest_message_id=88)]
        with patch.object(command,'_fetch_forum_topics',AsyncMock(side_effect=topics)) as fetch:
            rec=await command._inspect_group(client,chat,NS(top_message=NS(id=77)),job)
        fetch.assert_awaited_once()
        self.assertEqual(rec.visibility,'PRIVATE');self.assertTrue(rec.visibility_verified)
        normalized=snapshot([rec]).records[0]
        self.assertEqual(normalized.category,'private_groups')
        self.assertIn('Private topic',normalized.topics)
        self.assertIn('Latest message ID: 88',normalized.topics)
        self.assertIn('Latest message ID: 77',normalized.metadata)

    async def test_alias_discovered_in_existing_full_response_does_not_trigger_topics(self):
        full=NS(type=command.ChatType.FORUM,description=None,members_count=42,is_forum=True,
                username='public_alias',visibility='PUBLIC',visibility_verified=True,visibility_source='Active alias')
        chat=NS(id=-1007,type=command.ChatType.FORUM,title='Forum',username=None)
        with patch.object(command,'_get_chat_metadata',AsyncMock(return_value=full)),patch.object(command,'_fetch_forum_topics',AsyncMock()) as topics:
            rec=await command._inspect_group(AsyncMock(),chat,NS(top_message=None),command.InventoryJob('test',42))
        topics.assert_not_awaited();self.assertEqual(rec.visibility,'PUBLIC')


class ReportIntegrity(unittest.TestCase):
    def test_html_and_xlsx_use_exact_same_categories_counts_ids_and_metadata(self):
        from openpyxl import load_workbook
        s=snapshot()
        with TemporaryDirectory() as directory:
            html=Path(directory)/'report.html';xlsx=Path(directory)/'report.xlsx'
            render_html(s,html);render_xlsx(s,xlsx)
            markup=Markup(html.read_text())
            self.assertEqual([a['data-category'] for a in markup.articles],list(CATEGORY_IDS))
            self.assertEqual([a['data-filter'] for a in markup.stats],list(CATEGORY_IDS))
            self.assertTrue(set(CATEGORY_IDS).issubset({a.get('value') for a in markup.options}))
            self.assertIn('class="detailed"',html.read_text());self.assertIn('class="compact"',html.read_text())
            self.assertNotIn('<script>bad()</script>',html.read_text())
            from html import escape
            for r in s.records:
                if r.category in ('private_channels','private_groups'):
                    self.assertIn(escape(r.metadata,quote=True),html.read_text())
                    self.assertIn(escape(r.topics,quote=True),html.read_text())
            book=load_workbook(xlsx)
            try:
                self.assertEqual(book.sheetnames,['Summary','Private Channels','Private Groups','Bots',
                                                  'Private Chats','Public Channels','Public Groups','Unclassified'])
                for r in s.records:
                    sheet=book[CATEGORY_LABELS[r.category]]
                    self.assertEqual(sheet.max_row-1,s.counts[r.category])
                    compact=r.category in COMPACT_CATEGORIES
                    self.assertEqual(tuple(c.value for c in sheet[1]),COMPACT_HEADERS if compact else DETAILED_HEADERS)
                    self.assertEqual(sheet.cell(2,2 if compact else 3).value,r.chat_id)
                    self.assertEqual(sheet.cell(2,2 if compact else 3).data_type,'s')
                    self.assertFalse(any(c.data_type=='f' for row in sheet for c in row))
                    self.assertEqual(sheet.freeze_panes,'A2');self.assertEqual(sheet.auto_filter.ref,sheet.dimensions)
                    if r.category in ('private_channels','private_groups'):
                        self.assertEqual(sheet['I2'].value,r.metadata)
                        self.assertEqual(sheet['H2'].value,r.topics)
                        self.assertIn('Member count: 0',sheet['I2'].value)
                self.assertIn('Discussion group ID: -100000000000002',book['Private Channels']['I2'].value)
                self.assertIn('Has discussion: False',book['Private Groups']['I2'].value)
                self.assertIn('First message ID: 5',book['Private Groups']['H2'].value)
            finally:book.close()

    def test_long_private_topic_text_is_lossless_across_excel_continuations(self):
        from openpyxl import load_workbook
        row=fixture_records()[1];row.topics[0].title=('topic 😀  \n'*6000)+' final '
        s=snapshot([row])
        with TemporaryDirectory() as directory:
            path=Path(directory)/'a.xlsx';render_xlsx(s,path)
            book=load_workbook(path)
            try:
                sheet=book['Private Groups'];headers=[c.value for c in sheet[1]]
                text=sheet['H2'].value
                for part in range(2,100):
                    label=f'Topic Information (continued {part})'
                    if label not in headers:break
                    text+=sheet.cell(2,headers.index(label)+1).value
                self.assertEqual(text,s.records[0].topics)
                self.assertGreater(part,2)
                self.assertEqual(sheet.max_row,2)
            finally:book.close()

    def test_generated_artifact_corruption_is_detected(self):
        from openpyxl import load_workbook
        s=snapshot()
        with TemporaryDirectory() as directory:
            html=Path(directory)/'a.html';xlsx=Path(directory)/'a.xlsx'
            render_html(s,html);render_xlsx(s,xlsx)
            html.write_text(html.read_text().replace('data-category="private_channels"','data-category="bots"',1))
            with self.assertRaises(ReportValidationError):validate_html(s,html)
            book=load_workbook(xlsx);book['Private Groups'].delete_rows(2);book.save(xlsx);book.close()
            with self.assertRaises(ReportValidationError):validate_xlsx(s,xlsx)

    def test_empty_categories_and_missing_optional_values(self):
        for rows in ([],[command.DialogRecord(7,'unknown')]):
            with self.subTest(rows=len(rows)),TemporaryDirectory() as directory:
                s=snapshot(rows);render_html(s,Path(directory)/'a.html');render_xlsx(s,Path(directory)/'a.xlsx')
                self.assertEqual(sum(s.counts.values()),len(rows))

    def test_generated_javascript_all_categories_search_clear_and_empty_states(self):
        with TemporaryDirectory() as directory:
            path=Path(directory)/'a.html';render_html(snapshot(),path)
            text=path.read_text();markup=Markup(text)
        fixtures=[{'dataset':{'search':a['data-search'],'category':a['data-category'],'privacy':a['data-privacy']},'hidden':False} for a in markup.articles]
        script=re.search(r'<script>(.*?)</script>',text,re.S).group(1)
        harness="const assert=require('assert');const fixtures="+json.dumps(fixtures)+";const categories="+json.dumps(CATEGORY_IDS)+";"+'''
const elements={};for(const id of ['search','category','privacy','count','empty','clear'])elements[id]={value:'',addEventListener(){}};
const stats=categories.map(id=>({dataset:{filter:id}}));
global.document={querySelectorAll:s=>s==='article'?fixtures:s==='.stat'?stats:[],getElementById:id=>elements[id]};
'''+script+'''
for(const id of categories){category.value=id;filter();const visible=cards.filter(c=>!c.hidden);assert.equal(visible.length,1);assert.equal(visible[0].dataset.category,id);}
category.value='private_channels';search.value='ЎЗБЕК STRASSE';filter();assert.equal(cards.filter(c=>!c.hidden).length,1);
search.value='000000000001';filter();assert.equal(cards.filter(c=>!c.hidden).length,1);
search.value='missing';filter();assert.equal(elements.empty.hidden,false);
elements.clear.onclick();assert.equal(category.value,'private_channels');assert.equal(elements.empty.hidden,true);
privacy.value='PRIVATE';elements.clear.onclick();assert.equal(privacy.value,'PRIVATE');
privacy.value='';search.value='helper_bot';category.value='bots';filter();assert.equal(cards.filter(c=>!c.hidden).length,1);
search.value='';category.value='';filter();assert.equal(cards.filter(c=>!c.hidden).length,7);
stats[1].onclick();assert.equal(category.value,'private_groups');assert.equal(cards.filter(c=>!c.hidden).length,1);
'''
        result=subprocess.run(['node','-e',harness],capture_output=True,text=True,timeout=15)
        self.assertEqual(result.returncode,0,result.stderr)


if __name__=='__main__':unittest.main()
