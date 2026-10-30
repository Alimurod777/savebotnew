"""Isolated metadata export security and workflow regression tests."""
from __future__ import annotations
import asyncio
from contextlib import asynccontextmanager, contextmanager
from pathlib import Path
import tempfile
from types import SimpleNamespace as NS
import unittest
from unittest.mock import AsyncMock, patch
from core.userdata_reports import CATEGORY_IDS, CATEGORY_LABELS, CATEGORIES, COMPACT_CATEGORIES, normalize, render_html, render_xlsx


def source(records=()):
    from TechVJ.userdata_command import InventoryJob
    j = InventoryJob('abc123', 42, records=list(records), dialogs_discovered=len(records))
    j._cancel_event = asyncio.Event()
    return j


def record(**kw):
    from TechVJ.userdata_command import DialogRecord
    return DialogRecord(**dict(chat_id=-100123456789012345, category='channel', **kw))


class Reports(unittest.TestCase):
    def test_categories_unknown_dedup(self):
        r = record(title='unverified', visibility='PRIVATE')
        s = normalize(source([r, r]))
        self.assertEqual(len(s.records), 1)
        self.assertEqual(s.records[0].category, 'unclassified')
        self.assertEqual(tuple(s.counts), CATEGORY_IDS)

    def test_safe_html_unicode(self):
        s = normalize(source([record(title='<script>alert(1)</script> 😀 Ўзбек '+ 'x'*1000)]))
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)/'a.html'; render_html(s, p); text = p.read_text()
        self.assertNotIn('<script>alert(1)</script>', text)
        self.assertIn('&lt;script&gt;', text)
        self.assertIn('😀 Ўзбек', text)
        self.assertNotIn('innerHTML', text)
        self.assertIn("connect-src 'none'", text)

    def test_workbook_values_and_injection(self):
        from openpyxl import load_workbook
        s = normalize(source([record(title='=HYPERLINK("evil")', username='name')]))
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)/'a.xlsx'; render_xlsx(s,p); w=load_workbook(p)
            self.assertEqual(w.sheetnames, ['Summary', *CATEGORIES])
            sh=w['Public Channels']
            self.assertEqual(sh['A2'].data_type,'s')
            self.assertEqual(sh['B2'].value, '-100123456789012345')
            self.assertEqual(sh['B2'].data_type,'s')
            self.assertEqual(w['Summary']['B2'].value,1)
            self.assertEqual(sh.freeze_panes,'A2');w.close()

    def test_empty(self):
        s=normalize(source())
        with tempfile.TemporaryDirectory() as d:
            render_html(s,Path(d)/'a.html');render_xlsx(s,Path(d)/'a.xlsx')
        self.assertEqual(sum(s.counts.values()),0)


class Workflow(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        from TechVJ import userdata_command as m
        self.owner = patch.object(m, 'OWNER_ID', 42)
        self.owner.start()
        self.addCleanup(self.owner.stop)
        # Renderers are tested with real files above. Inline dispatch isolates the
        # workflow from this hosting runtime's hanging default-executor shutdown.
        async def dispatch(func, *args):
            return func(*args)
        self.dispatch = patch('asyncio.to_thread', dispatch)
        self.dispatch.start()
        self.addCleanup(self.dispatch.stop)

    async def test_target_progress_goes_only_to_owner(self):
        from TechVJ import userdata_command as m
        bot=AsyncMock();bot.send_message.return_value=NS(id=123)
        job=source(); job.target_user_id=99
        reporter=m.ProgressReporter(bot,job)
        await reporter.send_initial()
        await reporter.update(force=True)
        await reporter.finalize('Done')
        self.assertEqual(bot.send_message.call_args.args[0],42)
        self.assertTrue(all(c.args[0]==42 for c in bot.edit_message_text.call_args_list))

    async def test_duplicate_job_rejected(self):
        from TechVJ import userdata_command as m
        msg=NS(from_user=NS(id=42),chat=NS(id=42,type=m.ChatType.PRIVATE),text='/userdata 42',reply=AsyncMock())
        with patch.object(m,'_init_job_tables',AsyncMock()),patch.object(m,'_find_active_job_for_user',AsyncMock(return_value='abc')),patch.dict(m._active_jobs,{'abc':source()}),patch.object(m.async_db,'find_user',AsyncMock()) as lookup:
            await m.cmd_userdata(AsyncMock(),msg)
            lookup.assert_not_awaited()
            self.assertIn('already running',msg.reply.call_args.args[0])

    async def test_owner_selects_target_session(self):
        from TechVJ import userdata_command as m
        msg=NS(from_user=NS(id=42),chat=NS(id=42,type=m.ChatType.PRIVATE),text='/userdata 99',reply=AsyncMock())
        with patch.object(m,'_init_job_tables',AsyncMock()),patch.object(m,'_find_active_job_for_user',AsyncMock(return_value=None)),patch.object(m.async_db,'find_user',AsyncMock(return_value=None)) as lookup:
            await m.cmd_userdata(AsyncMock(),msg)
            lookup.assert_awaited_once_with(99)

    async def test_retry_creates_new_coroutine(self):
        from TechVJ import userdata_command as m
        call=AsyncMock(side_effect=[asyncio.TimeoutError(), 'ok'])
        with patch.object(m,'calculate_backoff',return_value=0):
            self.assertEqual(await m._safe_api_call(call), 'ok')
        self.assertEqual(call.await_count,2)

    async def test_nonowner_cannot_export_any_account(self):
        from TechVJ import userdata_command as m
        msg=NS(from_user=NS(id=7),chat=NS(id=7,type=m.ChatType.PRIVATE),text='/userdata 99',reply=AsyncMock())
        with patch.object(m.async_db,'find_user',new_callable=AsyncMock) as find:
            await m.cmd_userdata(AsyncMock(),msg);find.assert_not_awaited()
        msg.chat.id=-99
        self.assertFalse(await m._private_request(msg))

    async def test_nonowner_cannot_operate_jobs(self):
        from TechVJ import userdata_command as m
        msg=NS(from_user=NS(id=7),chat=NS(id=7,type=m.ChatType.PRIVATE),text='/userdata_status abc',reply=AsyncMock())
        for handler in (m.cmd_userdata_status,m.cmd_userdata_cancel,m.cmd_userdata_resume):
            with patch.object(m,'_init_job_tables',AsyncMock()) as init,patch.object(m,'_find_job_by_prefix',AsyncMock()) as lookup:
                await handler(AsyncMock(),msg)
            init.assert_not_awaited()
            lookup.assert_not_awaited()
        msg.reply.assert_not_awaited()

    async def test_identity_mismatch_stops_scan(self):
        from TechVJ import userdata_command as m
        client=AsyncMock();client.get_me.return_value=NS(id=99)
        @asynccontextmanager
        async def session(*a,**kw):yield client
        with patch('TechVJ.session_handler.create_user_session',session),patch.object(m,'_save_job',AsyncMock()):
            j=source();await m._run_inventory(AsyncMock(),j,'secret')
        client.get_dialogs.assert_not_called()
        self.assertEqual(j.status,m.JobStatus.FAILED)

    async def test_single_scan_shared_snapshot_and_owner_delivery(self):
        from TechVJ import userdata_command as m
        client=AsyncMock();client.get_me.return_value=NS(id=99)
        calls=[]
        async def dialogs():
            calls.append(1)
            yield NS(chat=NS(id=7,type=m.ChatType.PRIVATE,first_name='A',last_name='',username=None,is_bot=False),top_message=None)
        client.get_dialogs=dialogs
        @asynccontextmanager
        async def session(*a,**kw):yield client
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        snapshots=[];paths=[]
        def renderer(s,p):snapshots.append(s);Path(p).write_text('report');paths.append(p)
        with tempfile.TemporaryDirectory() as d,patch.object(m,'TEMP_DIR',d),patch('TechVJ.session_handler.create_user_session',session),patch.object(m,'_save_job',AsyncMock()),patch.object(m,'render_html',renderer),patch.object(m,'render_xlsx',renderer):
            j=source();j.target_user_id=99;await m._run_inventory(bot,j,'secret')
            self.assertEqual(calls,[1]);self.assertIs(snapshots[0],snapshots[1])
            self.assertEqual(snapshots[0].target_user_id,99)
            self.assertTrue(all(not Path(p).exists() for p in paths))
            self.assertEqual(bot.send_document.await_count,2)
            self.assertTrue(all(c.args[0]==42 for c in bot.send_document.call_args_list))
            self.assertTrue(all('99' in c.kwargs['caption'] for c in bot.send_document.call_args_list))
            self.assertEqual(j.status,m.JobStatus.COMPLETE)
        client.send_document.assert_not_called()

    async def test_partial_upload_cleanup_and_cancel(self):
        from TechVJ import userdata_command as m
        sender=AsyncMock();sender.get_me.return_value=NS(id=42)
        sender.send_document.side_effect=[RuntimeError('fail'),NS(id=1)]
        @asynccontextmanager
        async def session(*a,**kw):yield sender
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        bot.send_document.side_effect=[RuntimeError('fail'),NS(id=1)]
        def render(s,p):Path(p).write_text('report')
        with tempfile.TemporaryDirectory() as d,patch.object(m,'TEMP_DIR',d),patch('TechVJ.session_handler.create_user_session',session),patch.object(m,'render_html',render),patch.object(m,'render_xlsx',render):
            j=source();r=await m._deliver_reports(None,bot,j,normalize(j),AsyncMock(),'secret')
            self.assertEqual(r,{'HTML':False,'XLSX':True});self.assertEqual(list(Path(d).iterdir()),[])
            j._cancel_event.set()
            with self.assertRaises(asyncio.CancelledError):await m._deliver_reports(None,bot,j,normalize(j),AsyncMock(),'secret')
            self.assertEqual(list(Path(d).iterdir()),[])

    async def test_generation_failure_independent(self):
        from TechVJ import userdata_command as m
        sender=AsyncMock();sender.get_me.return_value=NS(id=42)
        @asynccontextmanager
        async def session(*a,**kw):yield sender
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        def fail(s,p):raise ValueError('bad')
        def good(s,p):Path(p).write_text('ok')
        with tempfile.TemporaryDirectory() as d,patch.object(m,'TEMP_DIR',d),patch('TechVJ.session_handler.create_user_session',session),patch.object(m,'render_html',fail),patch.object(m,'render_xlsx',good):
            j=source();r=await m._deliver_reports(None,bot,j,normalize(j),AsyncMock(),'secret')
            self.assertEqual(r,{'HTML':False,'XLSX':True});self.assertEqual(list(Path(d).iterdir()),[])

class OfflineControls(unittest.TestCase):
    def test_actual_javascript_filters(self):
        import json, re, subprocess
        s=normalize(source([record(title='Straße 😀',username='example')]))
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'a.html';render_html(s,p)
            js=re.search(r'<script>(.*?)</script>',p.read_text(),re.S).group(1)
            harness='''const assert=require('assert');
const fixtures=[{dataset:{search:'strasse 😀 -100123456789012345 example',category:'public_channels',privacy:'PUBLIC'},hidden:false}];
const elements={};for(const id of ['search','category','privacy','count','empty','clear'])elements[id]={value:'',addEventListener(){}};
global.document={querySelectorAll:s=>s==='article'?fixtures:[],getElementById:id=>elements[id]};
'''+js+'''
search.value='STRASSE';filter();assert.equal(cards[0].hidden,false);
search.value='123456789';filter();assert.equal(cards[0].hidden,false);
search.value='EXAMPLE';filter();assert.equal(cards[0].hidden,false);
privacy.value='PRIVATE';filter();assert.equal(cards[0].hidden,true);
elements.clear.onclick();assert.equal(privacy.value,'PRIVATE');assert.equal(cards[0].hidden,true);
privacy.value='';category.value='bots';filter();assert.equal(cards[0].hidden,true);assert.equal(elements.empty.hidden,false);
'''
            subprocess.run(['node','-e',harness],check=True,capture_output=True)


class OwnerCommands(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = Workflow.asyncSetUp

    def message(self, text, user_id=42):
        from TechVJ import userdata_command as m
        return NS(from_user=NS(id=user_id),chat=NS(id=user_id,type=m.ChatType.PRIVATE),text=text,reply=AsyncMock())

    async def test_all_commands_deny_nonowner_before_database_access(self):
        from TechVJ import userdata_command as m
        for handler, text in ((m.cmd_userdata,'/userdata'),(m.cmd_userdata,'/userdata 7'),
                              (m.cmd_userdata_status,'/userdata_status abc'),
                              (m.cmd_userdata_resume,'/userdata_resume abc'),
                              (m.cmd_userdata_cancel,'/userdata_cancel abc')):
            with self.subTest(text=text), patch.object(m,'_init_job_tables',AsyncMock()) as init, patch.object(m.async_db,'find_user',AsyncMock()) as lookup:
                msg=self.message(text,7)
                await handler(AsyncMock(),msg)
                init.assert_not_awaited();lookup.assert_not_awaited();msg.reply.assert_not_awaited()

    async def test_owner_invalid_targets_and_group_context_do_not_access_database(self):
        from TechVJ import userdata_command as m
        for text in ('/userdata','/userdata invalid','/userdata -1','/userdata 0'):
            with self.subTest(text=text),patch.object(m.async_db,'find_user',AsyncMock()) as lookup:
                msg=self.message(text)
                await m.cmd_userdata(AsyncMock(),msg)
                lookup.assert_not_awaited();msg.reply.assert_awaited_once()
        msg=self.message('/userdata 99');msg.chat.id=-99
        with patch.object(m.async_db,'find_user',AsyncMock()) as lookup:
            await m.cmd_userdata(AsyncMock(),msg)
            lookup.assert_not_awaited()

    async def test_owner_target_job_uses_owner_task_and_target_session(self):
        from TechVJ import userdata_command as m
        for resume in (False,True):
            j=source();j.target_user_id=99;j.status=m.JobStatus.FAILED
            msg=self.message('/userdata_resume abc' if resume else '/userdata 99')
            async def create_task(user_id,coro,**kw):
                try:
                    self.assertEqual(user_id,42)
                    self.assertIn(99,[job.target_user_id for job in m._active_jobs.values()])
                finally:
                    coro.close()
                return NS()
            with self.subTest(resume=resume),patch.dict(m._active_jobs,{},clear=True),patch.object(m,'_init_job_tables',AsyncMock()),patch.object(m,'_find_active_job_for_user',AsyncMock(return_value=None)),patch.object(m,'_find_job_by_prefix',AsyncMock(return_value=j)),patch.object(m,'_save_job',AsyncMock()),patch.object(m.async_db,'find_user',AsyncMock(return_value={'logged_in':True,'session':'test-only'})) as lookup,patch('TechVJ.task_manager.task_manager.create_task',new=AsyncMock(side_effect=create_task)) as launch:
                await (m.cmd_userdata_resume if resume else m.cmd_userdata)(AsyncMock(),msg)
                lookup.assert_awaited_once_with(99)
                launch.assert_awaited_once()

    async def test_owner_can_status_and_cancel_other_target_job(self):
        from TechVJ import userdata_command as m
        j=source();j.target_user_id=99
        with patch.object(m,'_init_job_tables',AsyncMock()),patch.object(m,'_find_job_by_prefix',AsyncMock(return_value=j)),patch.object(m,'_save_job',AsyncMock()) as save,patch.dict(m._active_jobs,{},clear=True):
            msg=self.message('/userdata_status abc')
            await m.cmd_userdata_status(AsyncMock(),msg)
            self.assertIn('Target user: 99',msg.reply.call_args.args[0])
            msg=self.message('/userdata_cancel abc')
            await m.cmd_userdata_cancel(AsyncMock(),msg)
            self.assertEqual(j.status,m.JobStatus.CANCELLED)
            save.assert_awaited_once_with(j)


class ReportAccuracy(unittest.TestCase):
    def test_reports_identify_target_account(self):
        from openpyxl import load_workbook
        j=source();j.target_user_id=99;s=normalize(j)
        with tempfile.TemporaryDirectory() as d:
            html_path=Path(d)/'a.html';xlsx_path=Path(d)/'a.xlsx'
            render_html(s,html_path);render_xlsx(s,xlsx_path)
            self.assertIn('Account: 99',html_path.read_text())
            book=load_workbook(xlsx_path)
            try:
                summary=dict(book['Summary'].values)
                self.assertEqual(summary['Account ID'],'99')
            finally:
                book.close()

    def test_order_counts_and_all_cells(self):
        from TechVJ.userdata_command import DialogRecord
        from openpyxl import load_workbook
        rows = [
            DialogRecord(1, 'bot', title='@bot'),
            DialogRecord(2, 'group', title='+public', username='public', chat_type='ChatType.SUPERGROUP'),
            DialogRecord(3, 'group', title='-private', visibility='PRIVATE', chat_type='ChatType.GROUP'),
            DialogRecord(4, 'private', title='=user', chat_type='ChatType.PRIVATE'),
            DialogRecord(5, 'channel', title='=public', username='channel'),
            DialogRecord(6, 'channel', title='private?', visibility='PRIVATE'),
        ]
        s = normalize(source(rows))
        self.assertEqual([r.category for r in s.records], [
            'private_groups', 'bots', 'private_chats', 'public_channels', 'public_groups', 'unclassified'])
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)/'a.xlsx'; render_xlsx(s, p)
            w = load_workbook(p)
            for i, (cat,label) in enumerate(CATEGORY_LABELS.items(), 5):
                self.assertEqual(w['Summary'].cell(i, 2).value, s.counts[cat])
                self.assertEqual(w[label].max_row - 1, s.counts[cat])
                expected = [r for r in s.records if r.category == cat]
                offset=0 if cat in COMPACT_CATEGORIES else 1
                for cells, r in zip(w[label].iter_rows(min_row=2), expected):
                    self.assertEqual(cells[offset].value, r.name)
                    self.assertEqual(cells[offset+1].value, r.chat_id)
                    self.assertTrue(all(c.data_type != 'f' for c in cells))
                self.assertEqual(w[label].auto_filter.ref, w[label].dimensions)
            w.close()

    def test_excel_preserves_long_detailed_fields_in_continuation_columns(self):
        from openpyxl import load_workbook
        s = normalize(source([record(title='x'*32768)]))
        with tempfile.TemporaryDirectory() as d:
            render_html(s, Path(d)/'a.html')
            render_xlsx(s, Path(d)/'a.xlsx')
            book=load_workbook(Path(d)/'a.xlsx')
            try:
                sheet=book['Unclassified'];headers=[cell.value for cell in sheet[1]]
                self.assertEqual(sheet['B2'].value+sheet.cell(2,headers.index('Chat Name (continued 2)')+1).value,'x'*32768)
            finally:
                book.close()

    def test_missing_unicode_and_csp_hash(self):
        import base64, hashlib, re
        s = normalize(source([record(title='bad\ud800\x00', description='a\x01b')]))
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)/'a.html'; render_html(s, p)
            text = p.read_text()
            script = re.search(r'<script>(.*?)</script>', text, re.S).group(1)
            digest = base64.b64encode(hashlib.sha256(script.encode()).digest()).decode()
            self.assertIn("script-src 'sha256-"+digest+"'", text)
            self.assertNotIn("script-src 'unsafe-inline'", text)
            render_xlsx(s, Path(d)/'a.xlsx')
        self.assertIn('Unknown / unverified', text)

    def test_topic_empty_known_versus_unavailable(self):
        a = record(title='known', is_forum=True, topic_count=0, topics_status='Complete accessible topic list')
        b = record(title='unknown', is_forum=True)
        from dataclasses import replace
        b = replace(b, chat_id=8)
        s = normalize(source([a, b]))
        self.assertIn('API reported count: 0', s.records[0].topics)
        self.assertIn('total unknown', s.records[1].topics)

    def test_type_aware_dedup_and_limit_disclosure(self):
        from dataclasses import replace
        a = record()
        j = source([a, replace(a, chat_type='different'), a])
        j.failures = [{'stage':'scan', 'reason':'Configured chat limit reached'}]
        s = normalize(j)
        self.assertEqual(len(s.records), 2)
        self.assertIn('lower bound', s.note)


class Reliability(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = Workflow.asyncSetUp

    async def test_both_formats_fail_independently_in_either_direction(self):
        from TechVJ import userdata_command as m
        sender=AsyncMock();sender.get_me.return_value=NS(id=42)
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        @asynccontextmanager
        async def session(*a, **kw): yield sender
        def good(s,p): Path(p).write_text('ok')
        def fail(s,p): raise RuntimeError('private error must not be disclosed')
        for failed_kind in ('HTML', 'XLSX'):
            for phase in ('generation', 'upload'):
                with self.subTest(kind=failed_kind, phase=phase):
                    bot.send_document.reset_mock()
                    bot.send_document.side_effect = ([RuntimeError(),NS(id=1)] if failed_kind=='HTML' else [NS(id=1),RuntimeError()]) if phase=='upload' else None
                    renderers = {'HTML':good, 'XLSX':good}
                    if phase=='generation': renderers[failed_kind]=fail
                    with tempfile.TemporaryDirectory() as d, patch.object(m,'TEMP_DIR',d), patch('TechVJ.session_handler.create_user_session',session), patch.object(m,'render_html',renderers['HTML']), patch.object(m,'render_xlsx',renderers['XLSX']):
                        j=source()
                        result=await m._deliver_reports(None,bot,j,normalize(j),AsyncMock(),'secret')
                        self.assertFalse(result[failed_kind])
                        self.assertTrue(result['XLSX' if failed_kind=='HTML' else 'HTML'])
                        self.assertNotIn('private error', str(j._report_errors))
                        self.assertEqual(list(Path(d).iterdir()),[])

    async def test_repeated_cancellation_joins_worker_before_cleanup(self):
        from TechVJ import userdata_command as m
        started, release = asyncio.Event(), asyncio.Event()
        async def dispatch(function, *args):
            started.set()
            await release.wait()
            return function(*args)
        def render(s,p): Path(p).write_text('ok')
        with tempfile.TemporaryDirectory() as d, patch.object(m,'TEMP_DIR',d), patch.object(m.asyncio,'to_thread',dispatch), patch.object(m,'render_html',render):
            j=source()
            task=asyncio.create_task(m._deliver_reports(None,AsyncMock(),j,normalize(j),AsyncMock(),'secret'))
            await started.wait()
            task.cancel()
            await asyncio.sleep(0)
            task.cancel()
            await asyncio.sleep(0)
            self.assertFalse(task.done())
            self.assertEqual(len(list(Path(d).iterdir())),1)
            release.set()
            with self.assertRaises(asyncio.CancelledError): await task
            self.assertEqual(list(Path(d).iterdir()),[])

    async def test_cleanup_failure_recorded_without_exposing_paths(self):
        from TechVJ import userdata_command as m
        j=source()
        with patch.object(m.tempfile, 'TemporaryDirectory') as directory:
            directory.return_value.name='/tmp/test-only'
            directory.return_value.cleanup.side_effect=PermissionError('secret path')
            with m._report_directory(j): pass
        self.assertIn('cleanup',j._report_errors)
        self.assertNotIn('secret',str(j._report_errors))

    async def test_long_floodwait_and_revocation_stop_retries(self):
        from TechVJ import userdata_command as m
        from TechVJ.session_handler import SessionInvalidError
        from wzgram.errors import SessionRevoked
        for error in (m.FloodWait(500), SessionRevoked()):
            api=AsyncMock(side_effect=error)
            with self.assertRaises((m.FloodWait, SessionInvalidError)):
                await m._safe_api_call(api)
            self.assertEqual(api.await_count,1)

    async def test_floodwait_is_honoured_outside_api_timeout_and_shown_in_progress(self):
        from TechVJ import userdata_command as m
        j=source();j.current_stage='Classifying channels'
        api=AsyncMock(side_effect=[m.FloodWait(30),'ok'])
        async def sleep(seconds):
            self.assertEqual(seconds,30)
            self.assertIn('waiting 30 seconds',j.current_stage)
        with patch.object(m.asyncio,'sleep',new=AsyncMock(side_effect=sleep)) as wait:
            self.assertEqual(await m._safe_api_call(api,job=j),'ok')
            wait.assert_awaited_once_with(30)
        self.assertEqual(api.await_count,2)
        self.assertEqual(j.current_stage,'Classifying channels')
        self.assertEqual(j._metadata_gap,2.0)

    async def test_cancel_event_during_flood_wait_prevents_retry(self):
        from TechVJ import userdata_command as m
        j=source();api=AsyncMock(side_effect=m.FloodWait(30))
        async def sleep(seconds):j._cancel_event.set()
        with patch.object(m.asyncio,'sleep',new=AsyncMock(side_effect=sleep)):
            with self.assertRaises(asyncio.CancelledError):
                await m._safe_api_call(api,job=j)
        self.assertEqual(api.await_count,1)

    async def test_metadata_requests_are_paced_and_disable_hidden_flood_sleeps(self):
        from TechVJ import userdata_command as m
        client=AsyncMock()
        client.resolve_peer.return_value=NS(channel_id=7,access_hash=9)
        client.invoke.return_value=NS(full_chat=NS(id=7),chats=[NS(id=7)])
        chat=NS(id=-1000000000007,type=m.ChatType.CHANNEL)
        j=source()
        with patch.object(m.time,'monotonic',return_value=100),patch.object(m.asyncio,'sleep',AsyncMock()) as sleep:
            await m._get_chat_metadata(client,chat,j)
            await m._get_chat_metadata(client,chat,j)
            sleep.assert_awaited_once_with(1.0)
        self.assertTrue(all(call.kwargs['sleep_threshold']==0 for call in client.invoke.call_args_list))
    async def test_delivery_to_owner_does_not_reopen_target_session(self):
        from TechVJ import userdata_command as m
        def render(s, p): Path(p).write_text('report')
        with tempfile.TemporaryDirectory() as d, patch.object(m, 'TEMP_DIR', d), patch('TechVJ.session_handler.create_user_session') as session, patch.object(m, 'render_html', render), patch.object(m, 'render_xlsx', render):
            j = source(); j.target_user_id=99; bot=AsyncMock()
            results=await m._deliver_reports(None, bot, j, normalize(j), AsyncMock(), 'secret')
            session.assert_not_called()
            self.assertTrue(all(results.values()))
            self.assertEqual(bot.send_document.await_count, 2)
            self.assertTrue(all(c.args[0] == 42 for c in bot.send_document.call_args_list))
            self.assertEqual(list(Path(d).iterdir()), [])

    async def test_upload_floodwait_then_transient_retry(self):
        from TechVJ import userdata_command as m
        sender = AsyncMock()
        sender.send_document.side_effect = [m.FloodWait(1), ConnectionError(), NS(id=1)]
        with patch.object(m.asyncio, 'sleep', AsyncMock()) as sleep:
            await m._upload_report(sender, 555, '/tmp/report', 'HTML')
            self.assertEqual(sender.send_document.await_count, 3)
            self.assertEqual(sleep.await_count, 2)

    async def test_progress_throttle_and_flood_cooldown(self):
        from TechVJ import userdata_command as m
        bot = AsyncMock(); bot.send_message.return_value = NS(id=1)
        reporter = m.ProgressReporter(bot, source())
        with patch.object(m.time, 'monotonic', return_value=100):
            await reporter.send_initial()
            await reporter.update()
            bot.edit_message_text.assert_not_called()
            bot.edit_message_text.side_effect = m.FloodWait(30)
            await reporter.update(force=True)
        with patch.object(m.time, 'monotonic', return_value=110):
            await reporter.update(force=True)
        self.assertEqual(bot.edit_message_text.await_count, 1)

    async def test_topic_short_page_continues_by_message_date(self):
        from TechVJ import userdata_command as m
        one = NS(id=1, title='one', top_message=10, date=100, closed=False, hidden=False)
        two = NS(id=2, title='two', top_message=20, date=200, closed=False, hidden=False)
        pages = [NS(count=2, topics=[one], order_by_create_date=False, messages=[NS(id=10, date=900)]),
                 NS(count=2, topics=[two], order_by_create_date=False, messages=[])]
        client = AsyncMock(); client.invoke.side_effect = pages
        rec = record()
        # Test protocol pagination independently of raw constructor boilerplate.
        with patch.object(m.types, 'ForumTopic', NS), patch.object(m.asyncio, 'sleep', AsyncMock()):
            topics = await m._fetch_forum_topics(client, 42, rec)
        self.assertEqual([t.topic_id for t in topics], [1,2])
        self.assertEqual(client.invoke.call_args_list[1].args[0].offset_date, 900)
        self.assertTrue(all(call.kwargs['sleep_threshold']==0 for call in client.invoke.call_args_list))
        self.assertEqual(rec.topic_count, 2)
        self.assertEqual(rec.topics_status, 'Complete accessible topic list')

    async def test_topic_failure_not_confirmed_empty(self):
        from TechVJ import userdata_command as m
        client = AsyncMock(); client.resolve_peer.side_effect = RuntimeError()
        rec = record()
        self.assertEqual(await m._fetch_forum_topics(client, 42, rec), [])
        self.assertIsNone(rec.topic_count)
        self.assertNotIn('Complete', rec.topics_status)

    async def test_metadata_has_no_implicit_pinned_message_request(self):
        from TechVJ import userdata_command as m
        client = AsyncMock()
        client.resolve_peer.return_value = NS(channel_id=7, access_hash=9)
        client.invoke.return_value = NS(full_chat=NS(id=7, about='Description', participants_count=12, pinned_msg_id=99),
                                        chats=[NS(id=7, username='test', forum=False)])
        meta = await m._get_chat_metadata(client, NS(id=-1000000000007, type=m.ChatType.CHANNEL))
        self.assertEqual(meta.description, 'Description')
        self.assertEqual(meta.members_count, 12)
        client.get_chat.assert_not_called()
        client.get_messages.assert_not_called()
        self.assertEqual(client.invoke.await_count, 1)
        self.assertEqual(client.invoke.call_args.kwargs['sleep_threshold'],0)

    async def test_report_limit_explained_and_no_upload(self):
        from TechVJ import userdata_command as m
        import os
        def render(s, p): Path(p).write_text('oversize')
        with tempfile.TemporaryDirectory() as d, patch.object(m, 'TEMP_DIR', d), patch.dict(os.environ, USERDATA_MAX_REPORT_BYTES='1'), patch.object(m, 'render_html', render), patch.object(m, 'render_xlsx', render):
            j = source(); bot = AsyncMock()
            result = await m._deliver_reports(None, bot, j, normalize(j), AsyncMock(), 'secret')
            self.assertFalse(any(result.values()))
            self.assertIn('size limit', j._report_errors['HTML'])
            bot.get_me.assert_not_called()
            self.assertEqual(list(Path(d).iterdir()), [])

    async def test_cancellation_during_upload_cleans_temp_files(self):
        from TechVJ import userdata_command as m
        sender = AsyncMock(); sender.get_me.return_value = NS(id=42)
        @asynccontextmanager
        async def session(*a, **kw): yield sender
        def render(s,p): Path(p).write_text('ok')
        paths=[]
        async def upload(chat, path, **kwargs):
            paths.append(path)
            self.assertTrue(Path(path).exists())
            raise asyncio.CancelledError()
        sender.send_document.side_effect=upload
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        bot.send_document.side_effect=upload
        with tempfile.TemporaryDirectory() as d, patch.object(m,'TEMP_DIR',d), patch('TechVJ.session_handler.create_user_session',session), patch.object(m,'render_html',render), patch.object(m,'render_xlsx',render):
            j=source()
            with self.assertRaises(asyncio.CancelledError):
                await m._deliver_reports(None,bot,j,normalize(j),AsyncMock(),'secret')
            self.assertTrue(paths)
            self.assertEqual(list(Path(d).iterdir()), [])

    async def test_inventory_stays_registered_until_cancel_cleanup_finishes(self):
        from TechVJ import userdata_command as m
        j=source(); client=AsyncMock();client.get_me.return_value=NS(id=42)
        async def dialogs():
            yield NS(chat=NS(id=7,type=m.ChatType.PRIVATE,first_name='A',last_name='',username=None,is_bot=False))
        client.get_dialogs=dialogs
        async def upload(*args, **kwargs):
            self.assertIn(j.job_id,m._active_jobs)
            raise asyncio.CancelledError()
        exits=[]
        @asynccontextmanager
        async def session(*a,**kw):
            yield client
        original_directory=m._report_directory
        @contextmanager
        def checked_directory(job):
            try:
                with original_directory(job) as directory:
                    yield directory
            finally:
                self.assertIn(j.job_id,m._active_jobs)
                self.assertEqual(list(Path(d).iterdir()),[])
                exits.append(True)
        def render(s,p):Path(p).write_text('ok')
        bot=AsyncMock();bot.get_me.return_value=NS(id=555)
        bot.send_document.side_effect=upload
        with tempfile.TemporaryDirectory() as d, patch.object(m,'TEMP_DIR',d), patch.dict(m._active_jobs,{j.job_id:j}), patch('TechVJ.session_handler.create_user_session',session), patch.object(m,'_report_directory',checked_directory), patch.object(m,'_save_job',AsyncMock()), patch.object(m,'render_html',render), patch.object(m,'render_xlsx',render):
            await m._run_inventory(bot,j,'secret')
            self.assertEqual(j.status,m.JobStatus.CANCELLED)
            self.assertEqual(exits,[True])
            self.assertNotIn(j.job_id,m._active_jobs)
            self.assertEqual(list(Path(d).iterdir()),[])


if __name__ == '__main__':
    unittest.main()
